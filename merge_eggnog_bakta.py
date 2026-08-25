#!/usr/bin/env python3
"""Add validated eggNOG-mapper evidence to pristine Bakta GenBank bytes."""

from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from merge_engine import (
    Insertion,
    MergeError,
    RawDocument,
    RawFeature,
    comment_insertion,
    finalize_merge,
    parse_faa,
    parse_genbank_bytes,
    qualifier_insertion,
    sha256_bytes,
    validate_faa_gbff,
    validate_genbank_semantics,
)

GO_RE = re.compile(r"GO:\d{7}\Z")
EC_RE = re.compile(r"\d+\.\d+\.\d+\.\d+\Z")
EC_PARTIAL_RE = re.compile(r"(?:\d+\.\d+\.\d+\.-|\d+\.\d+\.-\.-|\d+\.-\.-\.-)\Z")
KO_RE = re.compile(r"K\d{5}\Z")
COG_RE = re.compile(r"COG\d{4}\Z")
CAZY_RE = re.compile(r"(?:GH|GT|PL|CE|AA|CBM)\d+(?:_\d+)?\Z")
QUERY_RE = re.compile(r"(?![#*])\S+\Z")
ANNOTATION_TOKEN_RE = re.compile(r"[^\s,]+\Z")
VERSION_RE = re.compile(r"^##\s*emapper-([^\s]+)")
CONFIDENCE_RANK = {"low": 0, "medium": 1, "high": 2}
CONFIDENCE_CODE_RANK = {"l": 0, "m": 1, "h": 2}
CONFIDENCE_SCORED_FIELDS = (
    "Preferred_name",
    "GOs",
    "EC",
    "KEGG_ko",
    "KEGG_Pathway",
    "KEGG_Module",
    "KEGG_Reaction",
    "KEGG_rclass",
    "BRITE",
    "KEGG_TC",
    "CAZy",
    "BiGG_Reaction",
    "PFAMs",
)
CONFIDENCE_FIELD_INDEX = {
    field: index for index, field in enumerate(CONFIDENCE_SCORED_FIELDS)
}
CONFIDENCE_ORDER_PREFIX = "## confidence field order:"
REQUIRED_COLUMNS = (
    "#query",
    "seed_ortholog",
    "evalue",
    "score",
    "COG_category",
    "Preferred_name",
    "GOs",
    "EC",
    "KEGG_ko",
    "CAZy",
    "annotation_confidence",
)
COMMENT_MARKER = "##enrich-bakta:eggNOG:v1##"
FEATURE_NOTE_FIELDS = {"PFAMs", "eggNOG_OGs"}
CONTEXT_LIST_FIELDS = {
    "KEGG_Pathway",
    "KEGG_Module",
    "KEGG_Reaction",
    "KEGG_rclass",
    "BRITE",
    "KEGG_TC",
    "BiGG_Reaction",
}
HIGHER_ORDER_FIELDS = (
    "Description",
    "KEGG_Pathway",
    "KEGG_Module",
    "KEGG_Reaction",
    "KEGG_rclass",
    "BRITE",
    "KEGG_TC",
    "BiGG_Reaction",
    "tax_ceiling",
    "max_annot_lvl",
    "farthest_donor_taxid",
    "farthest_donor_lineage",
)
MISSING_ANNOTATION_VALUES = {"", "-", "NA"}


@dataclass(frozen=True)
class EggnogHit:
    row_number: int
    query_id: str
    seed_ortholog: str
    evalue: str
    score: str
    cog_category: str
    preferred_name: str
    gos: tuple[str, ...]
    ec: tuple[str, ...]
    ec_partial: tuple[str, ...]
    kegg_ko: tuple[str, ...]
    cazy: tuple[str, ...]
    pfams: tuple[str, ...]
    eggnog_ogs: tuple[str, ...]
    confidence: str
    raw_fields: tuple[tuple[str, str], ...]

    def raw(self, name: str) -> str:
        return dict(self.raw_fields).get(name, "")


@dataclass(frozen=True)
class EggnogTable:
    hits: tuple[EggnogHit, ...]
    columns: tuple[str, ...]
    version: str | None
    format: str
    confidence_field_order: tuple[str, ...]
    confidence_contract_source: str


def _decimal(value: str, field: str, row_number: int) -> Decimal:
    try:
        parsed = Decimal(value)
    except InvalidOperation as exc:
        raise MergeError(f"eggNOG row {row_number}: invalid {field} {value!r}") from exc
    if not parsed.is_finite() or (field == "E-value" and parsed < 0):
        raise MergeError(f"eggNOG row {row_number}: invalid {field} {value!r}")
    return parsed


def _tokens(
    value: str,
    *,
    field: str,
    row_number: int,
    pattern: re.Pattern[str],
    prefix: str = "",
) -> tuple[str, ...]:
    if value.strip() in MISSING_ANNOTATION_VALUES:
        return ()
    result: list[str] = []
    for raw in value.split(","):
        token = raw.strip()
        if prefix and token.lower().startswith(prefix.lower()):
            token = token[len(prefix) :]
        if not token or not pattern.fullmatch(token):
            raise MergeError(f"eggNOG row {row_number}: invalid {field} value {raw!r}")
        if token not in result:
            result.append(token)
    return tuple(result)


def _parse_cazy(value: str, row_number: int) -> tuple[str, ...]:
    if value.strip() in MISSING_ANNOTATION_VALUES:
        return ()
    result: list[str] = []
    for raw in value.split(","):
        family = raw.strip().split("|", 1)[0].strip()
        if not CAZY_RE.fullmatch(family):
            raise MergeError(f"eggNOG row {row_number}: invalid CAZy value {raw!r}")
        if family not in result:
            result.append(family)
    return tuple(result)


def _parse_ec(value: str, row_number: int) -> tuple[tuple[str, ...], tuple[str, ...]]:
    if value.strip() in MISSING_ANNOTATION_VALUES:
        return (), ()
    full: list[str] = []
    partial: list[str] = []
    for raw in value.split(","):
        token = raw.strip()
        if token.lower().startswith("ec:"):
            token = token[3:]
        if EC_RE.fullmatch(token):
            if token not in full:
                full.append(token)
        elif EC_PARTIAL_RE.fullmatch(token):
            if token not in partial:
                partial.append(token)
        else:
            raise MergeError(f"eggNOG row {row_number}: invalid EC value {raw!r}")
    return tuple(full), tuple(partial)


def _validate_confidence(value: str, row_number: int) -> str:
    if len(value) != 13 or any(char not in {"l", "m", "h", "-"} for char in value):
        raise MergeError(
            f"eggNOG row {row_number}: annotation_confidence must contain 13 l/m/h/- codes"
        )
    return value


def _build_hit(values: dict[str, str], row_number: int) -> EggnogHit:
    for column in REQUIRED_COLUMNS:
        if column not in values:
            raise MergeError(f"eggNOG table is missing required column {column!r}")
    query_id = values["#query"].strip()
    if not QUERY_RE.fullmatch(query_id):
        raise MergeError(
            f"eggNOG row {row_number}: invalid query identifier {query_id!r}"
        )
    seed = values["seed_ortholog"].strip()
    if not seed or seed == "-":
        raise MergeError(f"eggNOG row {row_number}: missing seed_ortholog")
    _decimal(values["evalue"].strip(), "E-value", row_number)
    _decimal(values["score"].strip(), "score", row_number)
    confidence = _validate_confidence(
        values["annotation_confidence"].strip(), row_number
    )
    cog = values["COG_category"].strip()
    if cog not in {"", "-"} and not (
        COG_RE.fullmatch(cog) or re.fullmatch(r"[A-Z]+", cog)
    ):
        raise MergeError(f"eggNOG row {row_number}: invalid COG_category {cog!r}")
    ec, ec_partial = _parse_ec(values["EC"].strip(), row_number)
    return EggnogHit(
        row_number=row_number,
        query_id=query_id,
        seed_ortholog=seed,
        evalue=values["evalue"].strip(),
        score=values["score"].strip(),
        cog_category=cog,
        preferred_name=values["Preferred_name"].strip(),
        gos=_tokens(
            values["GOs"].strip(), field="GO", row_number=row_number, pattern=GO_RE
        ),
        ec=ec,
        ec_partial=ec_partial,
        kegg_ko=_tokens(
            values["KEGG_ko"].strip(),
            field="KEGG_ko",
            row_number=row_number,
            pattern=KO_RE,
            prefix="ko:",
        ),
        cazy=_parse_cazy(values["CAZy"].strip(), row_number),
        pfams=_tokens(
            values.get("PFAMs", "").strip(),
            field="PFAMs",
            row_number=row_number,
            pattern=ANNOTATION_TOKEN_RE,
        ),
        eggnog_ogs=_tokens(
            values.get("eggNOG_OGs", "").strip(),
            field="eggNOG_OGs",
            row_number=row_number,
            pattern=ANNOTATION_TOKEN_RE,
        ),
        confidence=confidence,
        raw_fields=tuple(values.items()),
    )


def parse_eggnog_tsv(
    data: bytes, *, expected_version: str | None = None
) -> EggnogTable:
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise MergeError(f"eggNOG TSV is not UTF-8 at byte {exc.start}") from exc
    version: str | None = None
    declared_confidence_order: tuple[str, ...] | None = None
    columns: list[str] | None = None
    hits: list[EggnogHit] = []
    query_ids: set[str] = set()
    for row_number, line in enumerate(text.splitlines(), start=1):
        if not line:
            continue
        if line.startswith("##"):
            match = VERSION_RE.match(line)
            if match:
                version = match.group(1)
            if line.lower().startswith(CONFIDENCE_ORDER_PREFIX):
                order = tuple(line[len(CONFIDENCE_ORDER_PREFIX) :].strip().split())
                if declared_confidence_order is not None:
                    raise MergeError(
                        "eggNOG TSV contains duplicate confidence field-order legends"
                    )
                declared_confidence_order = order
            continue
        if line.startswith("#"):
            candidate = line.split("\t")
            if candidate[0] == "#query":
                if columns is not None:
                    raise MergeError("eggNOG TSV contains duplicate #query header")
                if len(candidate) != len(set(candidate)):
                    raise MergeError("eggNOG TSV header contains duplicate columns")
                columns = candidate
            elif columns is not None and "\t" in line:
                raise MergeError(
                    f"eggNOG row {row_number}: data row begins with reserved '#'; "
                    "query IDs must not start with '#'"
                )
            continue
        if columns is None:
            raise MergeError(
                f"eggNOG row {row_number}: data appears before #query header"
            )
        fields = line.split("\t")
        if len(fields) != len(columns):
            raise MergeError(
                f"eggNOG row {row_number}: expected {len(columns)} columns, found {len(fields)}"
            )
        hit = _build_hit(dict(zip(columns, fields, strict=True)), row_number)
        if hit.query_id in query_ids:
            raise MergeError(
                f"eggNOG row {row_number}: duplicate query ID {hit.query_id!r}"
            )
        query_ids.add(hit.query_id)
        hits.append(hit)
    if columns is None:
        raise MergeError("eggNOG TSV is missing the #query header")
    if not hits:
        raise MergeError("eggNOG TSV contains no annotation rows")
    confidence_order = declared_confidence_order or CONFIDENCE_SCORED_FIELDS
    if confidence_order != CONFIDENCE_SCORED_FIELDS:
        raise MergeError(
            "eggNOG confidence field order does not match the supported contract: "
            f"{confidence_order!r}"
        )
    header_order = tuple(
        column for column in columns if column in CONFIDENCE_FIELD_INDEX
    )
    expected_header_order = tuple(
        column for column in CONFIDENCE_SCORED_FIELDS if column in header_order
    )
    if header_order != expected_header_order:
        raise MergeError(
            "eggNOG scored columns conflict with the confidence field-order contract"
        )
    if expected_version is not None and version != expected_version:
        raise MergeError(
            f"eggNOG version mismatch: declared={version!r}, requested={expected_version!r}"
        )
    return EggnogTable(
        tuple(hits),
        tuple(columns),
        version,
        "tsv",
        confidence_order,
        "header_legend" if declared_confidence_order else "documented_default",
    )


def parse_eggnog_xlsx(
    path: Path, *, expected_version: str | None = None
) -> EggnogTable:
    if expected_version is None:
        raise MergeError(
            "--eggnog-version is required for XLSX input without TSV metadata"
        )
    try:
        import openpyxl  # type: ignore[import-untyped]
    except ImportError as exc:
        raise MergeError(
            "XLSX input requires optional dependency openpyxl; TSV needs no extra dependency"
        ) from exc
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    if workbook.sheetnames != ["annotations"]:
        raise MergeError(
            "eggNOG XLSX must contain exactly one sheet named 'annotations'"
        )
    worksheet = workbook["annotations"]
    rows = worksheet.iter_rows()
    try:
        header_cells = next(rows)
    except StopIteration as exc:
        raise MergeError("eggNOG XLSX is empty") from exc
    columns = [
        str(cell.value) if cell.value is not None else "" for cell in header_cells
    ]
    if len(columns) != len(set(columns)):
        raise MergeError("eggNOG XLSX header contains duplicate columns")
    hits: list[EggnogHit] = []
    query_ids: set[str] = set()
    for row_number, cells in enumerate(rows, start=2):
        if all(cell.value is None for cell in cells):
            continue
        if len(cells) != len(columns):
            raise MergeError(f"eggNOG XLSX row {row_number}: wrong cell count")
        values: list[str] = []
        for cell in cells:
            if cell.data_type == "f":
                raise MergeError(
                    f"eggNOG XLSX row {row_number}: formulas are not accepted"
                )
            values.append("" if cell.value is None else str(cell.value))
        hit = _build_hit(dict(zip(columns, values, strict=True)), row_number)
        if hit.query_id in query_ids:
            raise MergeError(
                f"eggNOG XLSX row {row_number}: duplicate query ID {hit.query_id!r}"
            )
        query_ids.add(hit.query_id)
        hits.append(hit)
    if not hits:
        raise MergeError("eggNOG XLSX contains no annotation rows")
    header_order = tuple(
        column for column in columns if column in CONFIDENCE_FIELD_INDEX
    )
    expected_header_order = tuple(
        column for column in CONFIDENCE_SCORED_FIELDS if column in header_order
    )
    if header_order != expected_header_order:
        raise MergeError(
            "eggNOG scored columns conflict with the confidence field-order contract"
        )
    return EggnogTable(
        tuple(hits),
        tuple(columns),
        expected_version,
        "xlsx",
        CONFIDENCE_SCORED_FIELDS,
        "documented_default",
    )


def parse_eggnog_path(
    path: Path, *, expected_version: str | None = None, data: bytes | None = None
) -> EggnogTable:
    if path.suffix.lower() == ".xlsx":
        return parse_eggnog_xlsx(path, expected_version=expected_version)
    return parse_eggnog_tsv(
        data if data is not None else path.read_bytes(),
        expected_version=expected_version,
    )


def _existing_values(feature: RawFeature, qualifier: str) -> set[str]:
    values = set(feature.values(qualifier))
    if qualifier == "EC_number":
        values.update(
            value[3:] for value in feature.values("db_xref") if value.startswith("EC:")
        )
    if qualifier == "db_xref":
        values.update(
            match.group(0)
            for note in feature.values("note")
            for match in re.finditer(
                r"(?:GO:\d{7}|KEGG:K\d{5}|CAZy:(?:GH|GT|PL|CE|AA|CBM)\d+(?:_\d+)?)",
                note,
            )
        )
    return values


def _index_paired_genes(
    base: RawDocument,
) -> dict[tuple[int, str, str], RawFeature | None]:
    index: dict[tuple[int, str, str], RawFeature | None] = {}
    for feature in base.features:
        if feature.feature_type != "gene" or not feature.locus_tag:
            continue
        key = (feature.record_index, feature.locus_tag, feature.location)
        index[key] = None if key in index else feature
    return index


def _paired_gene(
    index: dict[tuple[int, str, str], RawFeature | None], cds: RawFeature
) -> RawFeature | None:
    return index.get((cds.record_index, cds.locus_tag or "", cds.location))


def _normalized_gene(raw: str, clean_suffix: bool) -> str:
    value = raw.strip()
    return re.sub(r"_\d+\Z", "", value) if clean_suffix else value


def _confidence_passes(hit: EggnogHit, field: str, minimum: str) -> bool:
    index = CONFIDENCE_FIELD_INDEX.get(field)
    return index is None or (
        hit.confidence[index] != "-"
        and CONFIDENCE_CODE_RANK[hit.confidence[index]] >= CONFIDENCE_RANK[minimum]
    )


def _context_values(field: str, raw: str) -> tuple[str, ...]:
    value = raw.strip()
    if value in MISSING_ANNOTATION_VALUES:
        return ()
    if field not in CONTEXT_LIST_FIELDS:
        return (value,)
    return tuple(
        token.strip()
        for token in value.split(",")
        if token.strip() not in MISSING_ANNOTATION_VALUES
    )


def _confidence_status(hit: EggnogHit, field: str, minimum: str) -> str:
    index = CONFIDENCE_FIELD_INDEX.get(field)
    if index is None:
        return "not_scored"
    return (
        "passes"
        if _confidence_passes(hit, field, minimum)
        else "below_requested_threshold"
    )


def _collect_context(
    cds: dict[str, RawFeature],
    table: EggnogTable,
    *,
    min_confidence: str,
) -> tuple[list[dict[str, Any]], list[dict[str, Any]], int]:
    report_entries: list[dict[str, Any]] = []
    manifest_rows: list[dict[str, Any]] = []
    context_value_count = 0
    for hit in table.hits:
        feature = cds[hit.query_id]
        context: dict[str, dict[str, Any]] = {}
        for field in HIGHER_ORDER_FIELDS:
            raw = hit.raw(field)
            values = _context_values(field, raw)
            if not values:
                continue
            confidence_index = CONFIDENCE_FIELD_INDEX.get(field)
            confidence_code = (
                "" if confidence_index is None else hit.confidence[confidence_index]
            )
            confidence_status = _confidence_status(hit, field, min_confidence)
            context[field] = {
                "raw": raw,
                "values": list(values),
                "confidence_code": confidence_code,
                "confidence_status": confidence_status,
            }
            context_value_count += len(values)
            manifest_rows.append(
                {
                    "entry_type": "eggnog_context",
                    "evidence_class": "higher_order_context",
                    "row_number": hit.row_number,
                    "query_id": hit.query_id,
                    "record": feature.record_id,
                    "feature_type": feature.feature_type,
                    "locus_tag": feature.locus_tag or "",
                    "field": field,
                    "raw_value": raw,
                    "normalized_value": ", ".join(values),
                    "normalized_values": list(values),
                    "confidence_code": confidence_code,
                    "confidence_status": confidence_status,
                    "status": "sidecar_only",
                    "emitted_qualifiers": "",
                    "seed_ortholog": hit.seed_ortholog,
                    "e_value": hit.evalue,
                    "score": hit.score,
                }
            )
        if context:
            report_entries.append(
                {
                    "entry_type": "eggnog_context",
                    "row_number": hit.row_number,
                    "query_id": hit.query_id,
                    "record": feature.record_id,
                    "feature_type": feature.feature_type,
                    "locus_tag": feature.locus_tag or "",
                    "seed_ortholog": hit.seed_ortholog,
                    "e_value": hit.evalue,
                    "score": hit.score,
                    "context": context,
                }
            )
    return report_entries, manifest_rows, context_value_count


def plan_eggnog_additions(
    base: RawDocument,
    proteins: dict[str, str],
    table: EggnogTable,
    *,
    eggnog_data: bytes,
    faa_data: bytes,
    min_confidence: str = "low",
    clean_gene_suffix: bool = False,
    add_comment_note: bool = True,
    add_feature_provenance: bool = True,
    merge_timestamp: str | None = None,
    starting_order: int = 0,
) -> tuple[list[Insertion], list[dict[str, Any]], dict[str, Any]]:
    if min_confidence not in CONFIDENCE_RANK:
        raise MergeError(f"invalid eggNOG confidence threshold {min_confidence!r}")
    cds, validation = validate_faa_gbff(
        base, proteins, (hit.query_id for hit in table.hits), source_name="eggNOG"
    )
    insertions: list[Insertion] = []
    evidence: list[dict[str, Any]] = []
    order = starting_order
    planned = {
        (feature.start, qualifier): set(feature.values(qualifier))
        for feature in cds.values()
        for qualifier in ("gene", "db_xref", "EC_number", "note", "inference")
    }
    emitted_by_query: set[str] = set()
    paired_gene_index = _index_paired_genes(base)
    stats: dict[str, Any] = {
        **validation,
        "rows": len(table.hits),
        "format": table.format,
        "eggnog_version": table.version or "",
        "confidence_filtered": 0,
        "candidate_values": 0,
        "emitted_values": 0,
        "partial_ec_skipped": sum(len(hit.ec_partial) for hit in table.hits),
        "confidence_field_order": list(table.confidence_field_order),
        "confidence_contract_source": table.confidence_contract_source,
    }

    evidence.extend(
        {
            "entry_type": "eggnog_row",
            "row_number": hit.row_number,
            "query_id": hit.query_id,
            "raw_fields": dict(hit.raw_fields),
        }
        for hit in table.hits
    )
    for hit in table.hits:
        feature = cds[hit.query_id]
        paired = _paired_gene(paired_gene_index, feature)
        for partial in hit.ec_partial:
            evidence.append(
                {
                    "entry_type": "eggnog_candidate",
                    "evidence_class": "direct_annotation",
                    "row_number": hit.row_number,
                    "query_id": hit.query_id,
                    "record": feature.record_id,
                    "feature_type": feature.feature_type,
                    "locus_tag": feature.locus_tag or "",
                    "field": "EC",
                    "raw_value": hit.raw("EC"),
                    "source_token": partial,
                    "normalized_value": partial,
                    "confidence_code": hit.confidence[CONFIDENCE_FIELD_INDEX["EC"]],
                    "status": "skipped_partial_ec",
                    "emitted_qualifiers": "",
                    "seed_ortholog": hit.seed_ortholog,
                    "e_value": hit.evalue,
                    "score": hit.score,
                }
            )
        candidates: list[tuple[str, str, str, int | None]] = []
        if hit.preferred_name not in {"", "-"}:
            candidates.append(
                (
                    "Preferred_name",
                    "gene",
                    _normalized_gene(hit.preferred_name, clean_gene_suffix),
                    CONFIDENCE_FIELD_INDEX["Preferred_name"],
                )
            )
        candidates.extend(
            ("GOs", "db_xref", value, CONFIDENCE_FIELD_INDEX["GOs"])
            for value in hit.gos
        )
        candidates.extend(
            ("EC", "EC_number", value, CONFIDENCE_FIELD_INDEX["EC"]) for value in hit.ec
        )
        candidates.extend(
            ("KEGG_ko", "db_xref", f"KEGG:{value}", CONFIDENCE_FIELD_INDEX["KEGG_ko"])
            for value in hit.kegg_ko
        )
        if COG_RE.fullmatch(hit.cog_category):
            candidates.append(("COG_category", "note", f"COG:{hit.cog_category}", None))
        candidates.extend(
            ("CAZy", "db_xref", f"CAZy:{value}", CONFIDENCE_FIELD_INDEX["CAZy"])
            for value in hit.cazy
        )
        candidates.extend(
            ("PFAMs", "note", f"PFAM:{value}", CONFIDENCE_FIELD_INDEX["PFAMs"])
            for value in hit.pfams
        )
        candidates.extend(
            ("eggNOG_OGs", "note", f"eggNOG_OG:{value}", None)
            for value in hit.eggnog_ogs
        )
        for field, qualifier, value, confidence_index in candidates:
            stats["candidate_values"] += 1
            status = "existing"
            emitted: list[str] = []
            raw_value = hit.raw(field)
            if confidence_index is not None and not _confidence_passes(
                hit, field, min_confidence
            ):
                stats["confidence_filtered"] += 1
                status = "filtered_confidence"
            elif qualifier == "gene":
                feature_genes = [
                    item.strip() for item in feature.values("gene") if item.strip()
                ]
                paired_genes = (
                    [item.strip() for item in paired.values("gene") if item.strip()]
                    if paired
                    else ["missing paired gene"]
                )
                if feature_genes or paired_genes:
                    status = "existing_gene"
                elif not paired:
                    status = "unpaired_gene"
                else:
                    for target in (paired, feature):
                        insertions.append(
                            qualifier_insertion(
                                base.data,
                                target,
                                "gene",
                                value,
                                "eggNOG",
                                raw_value,
                                order,
                            )
                        )
                        order += 1
                        planned.setdefault((target.start, "gene"), set()).add(value)
                        emitted.append(f'{target.feature_type}:/gene="{value}"')
                    status = "emitted"
            elif value in planned[
                (feature.start, qualifier)
            ] or value in _existing_values(feature, qualifier):
                status = "existing"
            else:
                insertions.append(
                    qualifier_insertion(
                        base.data, feature, qualifier, value, "eggNOG", raw_value, order
                    )
                )
                order += 1
                planned[(feature.start, qualifier)].add(value)
                emitted.append(f'/{qualifier}="{value}"')
                status = "emitted"
            if emitted:
                emitted_by_query.add(hit.query_id)
                stats["emitted_values"] += len(emitted)
            evidence.append(
                {
                    "entry_type": "eggnog_candidate",
                    "evidence_class": (
                        "feature_note"
                        if field in FEATURE_NOTE_FIELDS
                        else "direct_annotation"
                    ),
                    "row_number": hit.row_number,
                    "query_id": hit.query_id,
                    "record": feature.record_id,
                    "feature_type": feature.feature_type,
                    "locus_tag": feature.locus_tag or "",
                    "field": field,
                    "raw_value": raw_value,
                    "source_token": (
                        value.removeprefix("PFAM:")
                        if field == "PFAMs"
                        else value.removeprefix("eggNOG_OG:")
                        if field == "eggNOG_OGs"
                        else value
                    ),
                    "normalized_value": value,
                    "confidence_code": ""
                    if confidence_index is None
                    else hit.confidence[confidence_index],
                    "status": status,
                    "emitted_qualifiers": " | ".join(emitted),
                    "seed_ortholog": hit.seed_ortholog,
                    "e_value": hit.evalue,
                    "score": hit.score,
                }
            )
        if hit.query_id in emitted_by_query and add_feature_provenance:
            inference = f"DESCRIPTION:similar to AA sequence:eggNOG:{hit.seed_ortholog}"
            values = planned[(feature.start, "inference")]
            if inference not in values:
                insertions.append(
                    qualifier_insertion(
                        base.data,
                        feature,
                        "inference",
                        inference,
                        "eggNOG provenance",
                        hit.seed_ortholog,
                        order,
                    )
                )
                order += 1
                values.add(inference)

    context_entries, context_rows, context_value_count = _collect_context(
        cds, table, min_confidence=min_confidence
    )
    evidence.extend(context_rows)
    eggnog_hash = sha256_bytes(eggnog_data)
    faa_hash = sha256_bytes(faa_data)
    if add_comment_note:
        lines = [
            f"Source eggNOG-mapper {table.version or 'unspecified'}; table sha256={eggnog_hash[:16]}",
            f"FAA sha256={faa_hash[:16]}; rows={len(table.hits)}, mapped CDSs={len(cds)}.",
            "eggNOG assignments are genomic evidence, not proof of activity or phenotype.",
        ]
        if merge_timestamp:
            lines.append(f"Merge timestamp: {merge_timestamp}")
        for record in base.records:
            insertion = comment_insertion(
                base,
                record,
                f"{COMMENT_MARKER}:{eggnog_hash[:12]}:{faa_hash[:12]}",
                lines,
                "eggNOG provenance",
                order,
            )
            if insertion is not None:
                insertions.append(insertion)
                order += 1
    stats.update(
        {
            "eggnog_sha256": eggnog_hash,
            "faa_sha256": faa_hash,
            "hit_cds": len(emitted_by_query),
            "pfam_candidates": sum(len(hit.pfams) for hit in table.hits),
            "eggnog_og_candidates": sum(len(hit.eggnog_ogs) for hit in table.hits),
            "context_hits": len(context_entries),
            "context_manifest_rows": len(context_rows),
            "context_values": context_value_count,
            "planned_insertions": len(insertions),
            "_context_report": {
                "schema": "enrich-bakta.eggnog-context.v1",
                "metadata": {
                    "operation": "eggnog-context-report",
                    "eggnog_version": table.version or "",
                    "eggnog_sha256": eggnog_hash,
                    "faa_sha256": faa_hash,
                    "min_eggnog_confidence": min_confidence,
                    "format": table.format,
                    "columns": list(table.columns),
                    "rows": len(table.hits),
                    "context_hits": len(context_entries),
                    "context_values": context_value_count,
                    "fields": list(HIGHER_ORDER_FIELDS),
                    "confidence_field_order": list(table.confidence_field_order),
                    "confidence_contract_source": table.confidence_contract_source,
                    "note": (
                        "Higher-order eggNOG context is reported separately from "
                        "feature-level qualifiers."
                    ),
                },
                "entries": context_entries,
            },
        }
    )
    return insertions, evidence, stats


def merge(
    bakta_path: Path,
    faa_path: Path,
    eggnog_path: Path,
    output_path: Path,
    *,
    manifest_path: Path | None = None,
    eggnog_version: str | None = None,
    min_confidence: str = "low",
    clean_gene_suffix: bool = False,
    add_comment_note: bool = True,
    add_feature_provenance: bool = True,
    merge_timestamp: str | None = None,
    context_report_path: Path | None = None,
) -> dict[str, Any]:
    base_data, faa_data, eggnog_data = (
        bakta_path.read_bytes(),
        faa_path.read_bytes(),
        eggnog_path.read_bytes(),
    )
    validate_genbank_semantics(base_data, "Bakta input")
    base = parse_genbank_bytes(base_data, "Bakta input")
    table = parse_eggnog_path(
        eggnog_path, expected_version=eggnog_version, data=eggnog_data
    )
    insertions, evidence, stats = plan_eggnog_additions(
        base,
        parse_faa(faa_data),
        table,
        eggnog_data=eggnog_data,
        faa_data=faa_data,
        min_confidence=min_confidence,
        clean_gene_suffix=clean_gene_suffix,
        add_comment_note=add_comment_note,
        add_feature_provenance=add_feature_provenance,
        merge_timestamp=merge_timestamp,
    )
    context_report = stats.pop("_context_report")
    context_report["metadata"]["operation"] = "eggnog-merge"
    final = finalize_merge(
        base_path=bakta_path,
        base_data=base_data,
        output_path=output_path,
        other_inputs=[faa_path, eggnog_path],
        insertions=insertions,
        manifest_path=manifest_path,
        evidence_rows=evidence,
        sidecar_path=context_report_path,
        sidecar_payload=context_report if context_report_path else None,
        metadata={
            "operation": "eggnog-merge",
            "eggnog_sha256": stats["eggnog_sha256"],
            "faa_sha256": stats["faa_sha256"],
            "eggnog_version": stats["eggnog_version"],
            "min_eggnog_confidence": min_confidence,
            "confidence_field_order": stats["confidence_field_order"],
            "confidence_contract_source": stats["confidence_contract_source"],
            "merge_timestamp": merge_timestamp or "",
        },
    )
    return {"eggnog": stats, **final}


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Merge sequence-validated eggNOG-mapper evidence onto a Bakta GBFF."
    )
    parser.add_argument("bakta", type=Path)
    parser.add_argument("faa", type=Path)
    parser.add_argument("eggnog", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--manifest", type=Path)
    parser.add_argument("--eggnog-version")
    parser.add_argument(
        "--min-eggnog-confidence", choices=tuple(CONFIDENCE_RANK), default="low"
    )
    parser.add_argument("--clean-gene-suffix", action="store_true")
    parser.add_argument("--no-comment-note", action="store_true")
    parser.add_argument("--no-feature-provenance", action="store_true")
    parser.add_argument("--merge-timestamp")
    parser.add_argument(
        "--context-report",
        type=Path,
        help="write higher-order eggNOG context as a deterministic JSON sidecar",
    )
    parser.add_argument("--json", action="store_true")
    args = parser.parse_args()
    for label, path in (
        ("Bakta", args.bakta),
        ("FAA", args.faa),
        ("eggNOG", args.eggnog),
    ):
        if not path.is_file():
            parser.error(f"{label} input not found: {path}")
    try:
        stats = merge(
            args.bakta,
            args.faa,
            args.eggnog,
            args.output,
            manifest_path=args.manifest,
            eggnog_version=args.eggnog_version,
            min_confidence=args.min_eggnog_confidence,
            clean_gene_suffix=args.clean_gene_suffix,
            add_comment_note=not args.no_comment_note,
            add_feature_provenance=not args.no_feature_provenance,
            merge_timestamp=args.merge_timestamp,
            context_report_path=args.context_report,
        )
    except MergeError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2
    if args.json:
        print(json.dumps(stats, indent=2, sort_keys=True))
    else:
        print("eggNOG merge passed FAA/GBFF identity and byte-preservation audits.")
        print(f"  annotation rows: {stats['eggnog']['rows']:,}")
        print(f"  emitted values:  {stats['eggnog']['emitted_values']:,}")
        print(f"  output SHA-256:  {stats['output_sha256']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
